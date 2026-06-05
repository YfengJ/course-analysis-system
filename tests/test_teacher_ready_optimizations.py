import json
import zipfile
import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from docx import Document
from openpyxl import Workbook

from app import create_app
from config import TestingConfig
from models import (
    AnalysisRevision,
    AnalysisSnapshot,
    Course,
    GraduationRequirement,
    ImportBatch,
    ObjectiveRequirementMap,
    Report,
    Student,
    db,
)
from services.course_insight_service import CourseInsightService
from services.auth_service import AuthService
from services.import_service import ImportService
from services.data_backup_service import DataBackupService
from services.course_archive_service import CourseArchiveService
from services.report_quality_service import ReportQualityService
from services.analysis_run_service import AnalysisRunService
from services.attainment_service import AttainmentService
from services.report_service import ReportService
from services.seed_service import DEFAULT_SEMESTER, create_generic_course_structure


class TeacherReadyOptimizationTest(unittest.TestCase):
    def setUp(self):
        self.app = create_app(TestingConfig)
        self.client = self.app.test_client()
        self.app_context = self.app.app_context()
        self.app_context.push()
        db.drop_all()
        db.create_all()

    def tearDown(self):
        db.session.remove()
        db.drop_all()
        self.app_context.pop()

    def _create_course(self, code="TEACHER001", name="教师使用测试课程"):
        course = Course(
            code=code,
            name=name,
            semester=DEFAULT_SEMESTER,
            department="计算机科学与技术系",
            major="示例专业",
            course_owner="测试教师",
            instructors="测试教师",
            expected_value=0.65,
            class_names="测试班",
        )
        db.session.add(course)
        db.session.flush()
        create_generic_course_structure(course)
        db.session.commit()
        return course

    def _build_score_file(self, path: Path, homework_score=4):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "成绩明细"
        sheet.append(["学号", "姓名", "班级", "课后作业", "大作业", "随堂测试", "期末考试", "上机实践"])
        sheet.append(["2026001", "张三", "测试班", homework_score, 4, 8, 60, 8])
        workbook.save(path)

    def _build_outline_upload(self):
        stream = BytesIO()
        document = Document()
        document.add_paragraph("示例教学大纲")
        document.save(stream)
        stream.seek(0)
        return stream

    def test_course_delete_removes_analysis_snapshots_and_revisions(self):
        course = self._create_course()
        db.session.add(
            AnalysisSnapshot(
                course_id=course.id,
                semester=DEFAULT_SEMESTER,
                class_scope="全部班级",
                version_no=1,
                student_count=0,
                summary_json=json.dumps({"ok": True}, ensure_ascii=False),
            )
        )
        db.session.add(
            AnalysisRevision(
                course_id=course.id,
                semester=DEFAULT_SEMESTER,
                class_scope="全部班级",
                qualitative_overrides_json="{}",
                analysis_note="教师修订说明",
            )
        )
        db.session.commit()

        response = self.client.post(f"/courses/{course.id}/delete", follow_redirects=True)

        self.assertEqual(response.status_code, 200)
        self.assertEqual(AnalysisSnapshot.query.filter_by(course_id=course.id).count(), 0)
        self.assertEqual(AnalysisRevision.query.filter_by(course_id=course.id).count(), 0)

    def test_report_download_rejects_report_from_another_course(self):
        course_a = self._create_course(code="A001", name="课程A")
        course_b = self._create_course(code="B001", name="课程B")
        report = Report(
            course_id=course_a.id,
            semester=DEFAULT_SEMESTER,
            class_scope="全部班级",
            word_path="/tmp/not-used.docx",
            html_snapshot="{}",
        )
        db.session.add(report)
        db.session.commit()

        response = self.client.get(f"/courses/{course_b.id}/reports/download/{report.id}")

        self.assertEqual(response.status_code, 404)

    def test_score_import_preview_does_not_write_students(self):
        course = self._create_course()
        with tempfile.TemporaryDirectory() as temp_dir:
            score_path = Path(temp_dir) / "成绩.xlsx"
            self._build_score_file(score_path)

            preview = ImportService.preview_score_files([score_path], course, DEFAULT_SEMESTER)

        self.assertTrue(preview["success"], preview.get("issues"))
        self.assertEqual(preview["imported_estimate"], 1)
        self.assertEqual(preview["files"][0]["record_count"], 1)
        self.assertIn("测试班", preview["files"][0]["classes"])
        self.assertEqual(Student.query.filter_by(course_id=course.id).count(), 0)
        self.assertEqual(ImportBatch.query.filter_by(course_id=course.id).count(), 0)

    def test_score_import_preview_rejects_scores_above_full_score_before_import(self):
        course = self._create_course()
        with tempfile.TemporaryDirectory() as temp_dir:
            score_path = Path(temp_dir) / "成绩异常.xlsx"
            self._build_score_file(score_path, homework_score=999)

            preview = ImportService.preview_score_files([score_path], course, DEFAULT_SEMESTER)
            result = ImportService.import_score_files([score_path], course, DEFAULT_SEMESTER)

        self.assertFalse(preview["success"])
        self.assertTrue(any("超出允许范围" in item for item in preview["issues"]))
        self.assertFalse(result["success"])
        self.assertEqual(Student.query.filter_by(course_id=course.id).count(), 0)

    def test_outline_preview_does_not_overwrite_course_until_confirmed(self):
        course = self._create_course(name="原课程名称")
        parsed_payload = {
            "course_code": "NEW001",
            "course_name": "预检课程名称",
            "objectives": [{"title": "课程目标1", "description": "预检目标"}],
            "assessment_support": [],
            "requirements": [],
        }

        with patch(
            "routes.course_routes.OutlineTemplateAdapter.extract",
            return_value={"payload": parsed_payload, "raw_text": "示例", "summary": "示例"},
        ), patch("routes.course_routes.ImportService.import_outline") as import_mock:
            response = self.client.post(
                f"/courses/{course.id}/outline",
                data={"action": "preview", "file": (self._build_outline_upload(), "outline.docx")},
                content_type="multipart/form-data",
            )

        self.assertEqual(response.status_code, 200)
        import_mock.assert_not_called()
        db.session.refresh(course)
        self.assertEqual(course.name, "原课程名称")
        with self.client.session_transaction() as sess:
            self.assertEqual(sess["pending_outline_import"]["course_id"], course.id)

    def test_report_archive_marks_final_version_without_changing_snapshot(self):
        course = self._create_course()
        for objective in course.objectives:
            objective.description = f"{objective.title}用于测试报告归档质量检查的正式目标描述。"
            requirement = GraduationRequirement(
                code=f"R{objective.sequence}",
                title="测试毕业要求",
                indicator_point=f"{objective.sequence}-1",
                description="用于报告归档质量检查的测试指标点。",
            )
            db.session.add(requirement)
            db.session.flush()
            db.session.add(
                ObjectiveRequirementMap(
                    objective_id=objective.id,
                    requirement_id=requirement.id,
                    support_strength="M",
                )
            )
        with tempfile.TemporaryDirectory() as temp_dir:
            score_path = Path(temp_dir) / "成绩.xlsx"
            self._build_score_file(score_path)
            result = ImportService.import_score_files([score_path], course, DEFAULT_SEMESTER)
            self.assertTrue(result["success"], result.get("issues"))
        summary = AttainmentService.calculate(course, DEFAULT_SEMESTER, "全部班级")
        AnalysisRunService.mark_complete(course.id, DEFAULT_SEMESTER, "全部班级", summary["student_count"], summary=summary)
        CourseInsightService.save_manual_for_scope(
            course.id,
            DEFAULT_SEMESTER,
            "全部班级",
            "教师确认本轮报告可以归档。",
            "下一轮继续跟踪薄弱目标。",
        )
        snapshot = {"total_quantitative_attainment": 0.81, "objective_results": []}
        report = Report(
            course_id=course.id,
            semester=DEFAULT_SEMESTER,
            class_scope="全部班级",
            html_snapshot=json.dumps(snapshot, ensure_ascii=False),
            summary_text="旧摘要",
        )
        db.session.add(report)
        db.session.commit()

        response = self.client.post(f"/courses/{course.id}/reports/{report.id}/archive", follow_redirects=True)

        self.assertEqual(response.status_code, 200)
        db.session.refresh(report)
        self.assertTrue(report.is_archived)
        self.assertIsNotNone(report.archived_at)
        self.assertEqual(json.loads(report.html_snapshot), snapshot)

    def test_manual_chapter_five_text_is_used_in_report_context(self):
        course = self._create_course()

        payload = CourseInsightService.save_manual_for_scope(
            course.id,
            DEFAULT_SEMESTER,
            "全部班级",
            overview_text="教师人工确认：本轮课程目标总体达成。",
            improvement_text="下轮将增加随堂练习和分层辅导。",
        )
        context = ReportService.build_report_context(course, DEFAULT_SEMESTER, "全部班级")

        self.assertEqual(payload["provider"], "人工编辑")
        self.assertTrue(context["chapter_five_ready"])
        self.assertEqual(context["chapter_five_source"], "manual_edit")
        self.assertIn("教师人工确认", context["chapter_five"]["overall_analysis"])
        self.assertIn("分层辅导", context["chapter_five"]["improvement_actions"][0]["action"])

    def test_report_quality_check_flags_missing_teacher_inputs(self):
        course = self._create_course()
        course.course_owner = "待完善"
        db.session.commit()

        result = ReportQualityService.check_course_report(course, DEFAULT_SEMESTER, "全部班级")

        self.assertFalse(result["ready"])
        self.assertGreaterEqual(result["blocking_count"], 1)
        self.assertTrue(any("课程负责人" in item["message"] for item in result["items"]))
        self.assertTrue(any("第五章" in item["message"] for item in result["items"]))

    def test_course_archive_package_contains_manifest_quality_and_report(self):
        course = self._create_course()
        with tempfile.TemporaryDirectory() as temp_dir:
            score_path = Path(temp_dir) / "成绩.xlsx"
            self._build_score_file(score_path)
            result = ImportService.import_score_files([score_path], course, DEFAULT_SEMESTER)
            self.assertTrue(result["success"], result.get("issues"))
            summary = AttainmentService.calculate(course, DEFAULT_SEMESTER, "全部班级")
            AnalysisRunService.mark_complete(course.id, DEFAULT_SEMESTER, "全部班级", summary["student_count"], summary=summary)
            CourseInsightService.save_manual_for_scope(
                course.id,
                DEFAULT_SEMESTER,
                "全部班级",
                "教师确认本轮达成情况稳定。",
                "下轮继续加强过程性评价。",
            )
            report, _ = ReportService.generate_word_report(course, DEFAULT_SEMESTER, "全部班级", temp_dir)

            archive_path = CourseArchiveService.build_archive(course, DEFAULT_SEMESTER, "全部班级", temp_dir)

            self.assertTrue(Path(archive_path).exists())
            with zipfile.ZipFile(archive_path) as package:
                names = set(package.namelist())
                self.assertIn("manifest.json", names)
                self.assertIn("analysis_summary.json", names)
                self.assertIn("quality_check.json", names)
                self.assertTrue(any(name.endswith(Path(report.word_path).name) for name in names))


class FileDatabaseConfig(TestingConfig):
    SQLALCHEMY_DATABASE_URI = ""
    BACKUP_FOLDER = ""


class BackupRestoreWorkflowTest(unittest.TestCase):
    def test_database_backup_and_restore_round_trip(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            db_path = Path(temp_dir) / "attainment_system.db"
            backup_dir = Path(temp_dir) / "backups"

            class RuntimeBackupConfig(FileDatabaseConfig):
                SQLALCHEMY_DATABASE_URI = f"sqlite:///{db_path}"
                BACKUP_FOLDER = str(backup_dir)
                UPLOAD_FOLDER = str(Path(temp_dir) / "uploads")
                EXPORT_FOLDER = str(Path(temp_dir) / "exports")
                REPORT_FOLDER = str(Path(temp_dir) / "exports" / "reports")
                SAMPLE_DATA_FOLDER = str(Path(temp_dir) / "sample_data")

            app = create_app(RuntimeBackupConfig)
            with app.app_context():
                db.drop_all()
                db.create_all()
                course = Course(code="BACKUP001", name="备份前课程", semester=DEFAULT_SEMESTER)
                db.session.add(course)
                db.session.commit()

                backup_path = DataBackupService.create_backup(app)
                course.name = "备份后课程"
                db.session.commit()

                DataBackupService.restore_backup(app, backup_path)
                restored = Course.query.filter_by(code="BACKUP001").first()

                self.assertIsNotNone(restored)
                self.assertEqual(restored.name, "备份前课程")
                with zipfile.ZipFile(backup_path) as package:
                    self.assertIn("manifest.json", package.namelist())
                    self.assertIn("database/attainment_system.db", package.namelist())


class AuthConfig(TestingConfig):
    LOGIN_DISABLED = False
    DEFAULT_ADMIN_USERNAME = "admin"
    DEFAULT_ADMIN_PASSWORD = "testpass"
    DEFAULT_ADMIN_DISPLAY_NAME = "测试管理员"


class AuthenticationWorkflowTest(unittest.TestCase):
    def setUp(self):
        self.app = create_app(AuthConfig)
        self.client = self.app.test_client()
        self.app_context = self.app.app_context()
        self.app_context.push()
        db.drop_all()
        db.create_all()
        AuthService.ensure_default_user()

    def tearDown(self):
        db.session.remove()
        db.drop_all()
        self.app_context.pop()

    def test_login_is_required_and_default_admin_can_enter(self):
        response = self.client.get("/courses/")
        self.assertEqual(response.status_code, 302)
        self.assertIn("/login", response.location)

        login_response = self.client.post(
            "/login",
            data={"username": "admin", "password": "testpass"},
            follow_redirects=True,
        )

        self.assertEqual(login_response.status_code, 200)
        self.assertIn("课程工作台".encode("utf-8"), login_response.data)

    def test_default_password_login_shows_change_password_hint(self):
        response = self.client.post(
            "/login",
            data={"username": "admin", "password": "testpass"},
            follow_redirects=True,
        )

        self.assertEqual(response.status_code, 200)
        self.assertIn("当前仍在使用初始密码".encode("utf-8"), response.data)
        self.assertIn("/account/password".encode("utf-8"), response.data)

    def test_user_can_change_password_and_old_password_stops_working(self):
        self.client.post("/login", data={"username": "admin", "password": "testpass"})

        change_response = self.client.post(
            "/account/password",
            data={
                "current_password": "testpass",
                "new_password": "new-strong-pass",
                "confirm_password": "new-strong-pass",
            },
            follow_redirects=True,
        )

        self.assertEqual(change_response.status_code, 200)
        self.assertIn("密码已更新".encode("utf-8"), change_response.data)

        self.client.post("/logout")
        old_login = self.client.post(
            "/login",
            data={"username": "admin", "password": "testpass"},
            follow_redirects=True,
        )
        self.assertIn("账号或密码不正确".encode("utf-8"), old_login.data)

        new_login = self.client.post(
            "/login",
            data={"username": "admin", "password": "new-strong-pass"},
            follow_redirects=True,
        )
        self.assertIn("课程工作台".encode("utf-8"), new_login.data)
        self.assertNotIn("当前仍在使用初始密码".encode("utf-8"), new_login.data)

    def test_change_password_requires_current_password(self):
        self.client.post("/login", data={"username": "admin", "password": "testpass"})

        response = self.client.post(
            "/account/password",
            data={
                "current_password": "wrong",
                "new_password": "new-strong-pass",
                "confirm_password": "new-strong-pass",
            },
            follow_redirects=True,
        )

        self.assertEqual(response.status_code, 200)
        self.assertIn("当前密码不正确".encode("utf-8"), response.data)
        self.assertIsNotNone(AuthService.authenticate("admin", "testpass"))


if __name__ == "__main__":
    unittest.main()
